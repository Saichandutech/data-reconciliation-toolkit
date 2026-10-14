"""Write reconciliation + validation results to a formatted Excel workbook and an HTML summary."""
from __future__ import annotations

import html
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .reconcile import ReconResult
from .validate import RuleResult, summarise

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
FAIL_FILL = PatternFill("solid", fgColor="FDE2E1")
PASS_FILL = PatternFill("solid", fgColor="DCFCE7")


def _format_sheet(ws, status_col: str | None = None) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left")
    for i, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    if status_col:
        headers = [c.value for c in ws[1]]
        idx = headers.index(status_col) + 1
        for row in ws.iter_rows(min_row=2):
            cell = row[idx - 1]
            cell.fill = FAIL_FILL if cell.value == "FAIL" else PASS_FILL


def write_excel(path: Path, title: str, recon: ReconResult, rules: list[RuleResult], source_rows: int) -> Path:
    val_summary = summarise(rules, source_rows)
    failures = [r.failed_rows.assign(failed_rule=r.name) for r in rules if r.failures]
    val_detail = pd.concat(failures, ignore_index=True) if failures else pd.DataFrame({"result": ["No failures"]})

    sheets = {
        "Summary": recon.summary(),
        "Control Totals": recon.control_totals,
        "Validation": val_summary,
        "Validation Failures": val_detail,
        "Field Mismatches": recon.mismatches,
        "Missing in Target": recon.missing_in_target,
        "Missing in Source": recon.missing_in_source,
        "Duplicates (Target)": recon.duplicates_target,
        "Duplicates (Source)": recon.duplicates_source,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            out = df if len(df) else pd.DataFrame({"result": ["None found"]})
            out.to_excel(xw, sheet_name=name, index=False)
            _format_sheet(xw.sheets[name], status_col="status" if name == "Validation" else None)
        ws = xw.sheets["Summary"]
        ws.insert_rows(1, 2)
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M}"
        ws.freeze_panes = None
    return path


def _table(df: pd.DataFrame, max_rows: int = 25) -> str:
    if df.empty:
        return "<p class='muted'>None found.</p>"
    extra = f"<p class='muted'>Showing {max_rows} of {len(df)} rows – see the Excel report for all.</p>" if len(df) > max_rows else ""
    return df.head(max_rows).to_html(index=False, border=0, classes="t", na_rep="—", escape=True) + extra


def write_html(path: Path, title: str, recon: ReconResult, rules: list[RuleResult], source_rows: int) -> Path:
    val = summarise(rules, source_rows)
    s = recon.summary().set_index("metric")["value"]
    failed_rules = int((val["status"] == "FAIL").sum())

    def card(label, value, bad=False):
        return f"<div class='card{' bad' if bad else ''}'><div class='v'>{value}</div><div class='l'>{html.escape(label)}</div></div>"

    cards = "".join(
        [
            card("Fully reconciled", f"{s['Fully reconciled (% of source)']}%"),
            card("Missing in target", int(s["Missing in target"]), s["Missing in target"] > 0),
            card("Missing in source", int(s["Missing in source"]), s["Missing in source"] > 0),
            card("Records with mismatches", int(s["Records with field mismatches"]), s["Records with field mismatches"] > 0),
            card("Duplicate keys (target)", int(s["Duplicate keys in target"]), s["Duplicate keys in target"] > 0),
            card("Validation rules failed", f"{failed_rules}/{len(val)}", failed_rules > 0),
        ]
    )
    by_col = (
        recon.mismatches.groupby("column").size().rename("mismatches").reset_index()
        if len(recon.mismatches) else pd.DataFrame()
    )
    notes = "".join(f"<li>{html.escape(n)}</li>" for n in recon.notes)

    doc = f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title><style>
body{{font-family:system-ui,-apple-system,Segoe UI,sans-serif;margin:0 auto;max-width:1000px;padding:24px 16px;color:#1f2937;background:#fff}}
h1{{font-size:24px;margin-bottom:4px}} h2{{font-size:17px;margin-top:32px;border-bottom:1px solid #e5e7eb;padding-bottom:6px}}
.muted{{color:#6b7280;font-size:13px}} .cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin-top:20px}}
.card{{border:1px solid #e5e7eb;border-radius:8px;padding:14px;background:#f9fafb}} .card.bad{{background:#fef2f2;border-color:#fecaca}}
.v{{font-size:24px;font-weight:700}} .l{{font-size:13px;color:#4b5563}}
.wrap{{overflow-x:auto}} table.t{{border-collapse:collapse;font-size:13px;width:100%}}
.t th{{background:#1f3a5f;color:#fff;text-align:left;padding:6px 8px}} .t td{{padding:6px 8px;border-bottom:1px solid #e5e7eb}}
</style></head><body>
<h1>{html.escape(title)}</h1><div class="muted">Generated {datetime.now():%Y-%m-%d %H:%M} · key: <code>{html.escape(recon.key)}</code></div>
<div class="cards">{cards}</div>
{f"<ul class='muted'>{notes}</ul>" if notes else ""}
<h2>Control totals</h2><div class="wrap">{_table(recon.control_totals)}</div>
<h2>Validation rules (source)</h2><div class="wrap">{_table(val)}</div>
<h2>Mismatches by field</h2><div class="wrap">{_table(by_col)}</div>
<h2>Field-level mismatches</h2><div class="wrap">{_table(recon.mismatches)}</div>
<h2>Missing in target</h2><div class="wrap">{_table(recon.missing_in_target)}</div>
<h2>Missing in source</h2><div class="wrap">{_table(recon.missing_in_source)}</div>
<h2>Duplicate keys in target</h2><div class="wrap">{_table(recon.duplicates_target)}</div>
</body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(doc, encoding="utf-8")
    return path
